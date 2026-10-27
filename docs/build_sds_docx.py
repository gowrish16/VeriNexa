import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = docx.Document()

    # Set page margins: 1 inch (72pt)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(5)

    # Header and Footer
    section = doc.sections[0]
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = f_p.add_run("VeriNexa SDS | Dept of MCA, AJIET Mangaluru")
    f_run.font.size = Pt(9)
    f_run.font.color.rgb = RGBColor(128, 128, 128)

    i = 0
    n = len(lines)

    in_code_block = False
    code_block_lines = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Handle code/diagram blocks
        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_block_lines = []
            else:
                in_code_block = False
                # Render code/diagram block
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F4F6F8")
                set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

                # Set subtle single border
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(f'''
                    <w:tcBorders {nsdecls("w")}>
                        <w:top w:val="single" w:sz="6" w:space="0" w:color="0D9488"/>
                        <w:left w:val="single" w:sz="18" w:space="0" w:color="0D9488"/>
                        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="0D9488"/>
                        <w:right w:val="single" w:sz="6" w:space="0" w:color="0D9488"/>
                    </w:tcBorders>
                ''')
                tcPr.append(tcBorders)

                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(0)
                cp.paragraph_format.line_spacing = 1.0

                code_text = "\n".join(code_block_lines)
                c_run = cp.add_run(code_text)
                c_run.font.name = 'Consolas'
                c_run.font.size = Pt(8.5)
                c_run.font.color.rgb = RGBColor(24, 43, 56)

                # Add a tiny space after table
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_after = Pt(4)
                p_spacer.paragraph_format.space_before = Pt(0)

            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line.rstrip("\r\n"))
            i += 1
            continue

        # Handle Markdown Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            # Parse table
            rows_data = []
            for tl in table_lines:
                # Skip divider line like |:---|:---|
                if re.match(r'^\|[\s\:\-]+(?:\|[\s\:\-]+)*\|$', tl):
                    continue
                cells = [c.strip() for c in tl.split("|")[1:-1]]
                rows_data.append(cells)

            if rows_data:
                col_count = max(len(r) for r in rows_data)
                tbl = doc.add_table(rows=len(rows_data), cols=col_count)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, "B0BEC5")

                for r_idx, row in enumerate(rows_data):
                    is_header = (r_idx == 0)
                    for c_idx, val in enumerate(row):
                        if c_idx < col_count:
                            cell = tbl.cell(r_idx, c_idx)
                            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.line_spacing = 1.05

                            # Format text (bold, math)
                            cleaned_val = re.sub(r'[\$\*]', '', val)
                            run = p.add_run(cleaned_val)
                            run.font.name = 'Calibri'

                            if is_header:
                                set_cell_background(cell, "16323B")
                                run.font.bold = True
                                run.font.size = Pt(9.5)
                                run.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "FFFFFF")
                                else:
                                    set_cell_background(cell, "F7F9FA")
                                run.font.size = Pt(9)
                                run.font.color.rgb = RGBColor(33, 37, 41)

                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.space_after = Pt(4)
            continue

        # Skip horizontal dividers
        if stripped == "---":
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            h_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(14, 42, 56) # Deep Slate/Teal
            i += 1
            continue

        if stripped.startswith("## "):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(5)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(13, 148, 136) # Teal
            i += 1
            continue

        if stripped.startswith("### "):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 70)
            i += 1
            continue

        # Bullet lists
        if stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            # Format bold within bullet
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Numbered list
        if re.match(r'^\d+\.\s+', stripped):
            match = re.match(r'^\d+\.\s+(.*)$', stripped)
            num_text = match.group(1) if match else stripped
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*)', num_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Regular paragraphs
        if stripped != "":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
        
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated DOCX at: {docx_path}")

if __name__ == "__main__":
    md_file = r"C:\Projects\VeriNexa\docs\VeriNexa_Software_Design_Specification.md"
    docx_file = r"C:\Projects\VeriNexa\docs\VeriNexa_Software_Design_Specification.docx"
    build_docx(md_file, docx_file)
